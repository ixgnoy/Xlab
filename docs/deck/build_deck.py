"""Build docs/deck/PersonaLab.pptx (TOKEN2049 Origins submission deck).

Re-run after updating EVIDENCE below:
    uv run --with python-pptx python docs/deck/build_deck.py

Rule: never put a guessed ID or tx hash in EVIDENCE. Leave "TODO: ..." until it is verified.
"""

from pathlib import Path

import re

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).with_name("PersonaLab.pptx")
EXPLORER = "https://preprod.cardanoscan.io/transaction/"

# ---------------------------------------------------------------- evidence (edit here)
EVIDENCE = {
    "coworker_id": "01a1156a-dbc8-7360-813b-043c0e148296",
    "vendor": "CardanoFish 01a1156a-7184-77bb-a15f-f712b72381be",
    "agent_identifier": (
        "67ab0c92c4ac1610895a1c965ee50aba41a8f1513b15240723b3bd0b10648ac9"
        "9be37c280788d2c1a9270b4960d342b1d4fb4a830b5beba022000000"
    ),
    "registration_tx": "ca5832b506f34a834a732eba2ff3f00cc52f2fb6bcf84016670deadc6af85dc1",
    "task_personal": "01a115be-d0ca-77d5-bf90-b98aa4487fec",
    "task_event": "01a115c6-b98d-75bd-b5b0-7ab2f85d4ab4",
    "task_paid": "01a115db-7a10-7148-8526-d516924c360c",
    "paid_status": "Escrow FundsLocked, collection pending",
    "escrow_tx": "TODO: escrow tx",
    "collection_tx": "TODO: collection tx",
    "net_receipt": "TODO: net tUSDM received (settlement.mjs)",
    "repo": "https://github.com/ixgnoy/personalab-coworker",
    "team": "TODO: team members",
}

# ---------------------------------------------------------------- theme
BG = RGBColor(0x0B, 0x0B, 0x10)
CARD = RGBColor(0x17, 0x17, 0x20)
CARD_LINE = RGBColor(0x2A, 0x2A, 0x38)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xA0, 0xA3, 0xB1)
ACCENT = RGBColor(0x8B, 0x5C, 0xF6)  # single accent (violet)
ACCENT_SOFT = RGBColor(0x24, 0x1B, 0x3D)
FONT = "Segoe UI"
MONO = "Consolas"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def is_todo(v: str) -> bool:
    return v.startswith("TODO")


def style_run(run, size=18, color=WHITE, bold=False, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def textbox(slide, x, y, w, h, text="", size=18, color=WHITE, bold=False, font=FONT,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        style_run(p.add_run(), size, color, bold, font)
        p.runs[0].text = line
    return tb


def rect(slide, x, y, w, h, fill=CARD, line=CARD_LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def box_text(shape, lines, size=14, color=WHITE, bold_first=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE, first_color=None):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right"):
        setattr(tf, m, Inches(0.12))
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        first = i == 0 and bold_first
        style_run(r, size if first else size - 2, (first_color or color) if first else MUTED, first)


def new_slide(title, kicker=None, notes=""):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    # accent bar
    rect(s, Inches(0.6), Inches(0.55), Inches(0.08), Inches(0.62), fill=ACCENT, line=None,
         shape=MSO_SHAPE.RECTANGLE)
    if kicker:
        textbox(s, Inches(0.85), Inches(0.38), Inches(10), Inches(0.35), kicker.upper(), 12, ACCENT, True)
    t = textbox(s, Inches(0.85), Inches(0.62), Inches(11.8), Inches(0.7), title, 32, WHITE, True)
    t.name = "Title"
    footer = textbox(s, Inches(0.6), SH - Inches(0.45), Inches(8), Inches(0.3),
                     "PersonaLab  ·  TOKEN2049 Origins  ·  Masumi + Sokosumi on Cardano Preprod", 10, MUTED)
    footer.name = "Footer"
    num = textbox(s, SW - Inches(1.2), SH - Inches(0.45), Inches(0.6), Inches(0.3),
                  str(len(prs.slides)), 10, MUTED, align=PP_ALIGN.RIGHT)
    num.name = "Number"
    if notes:
        s.notes_slide.notes_text_frame.text = notes.strip()
    return s


def bullets(slide, x, y, w, h, items, size=20, gap=10):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        head, _, rest = item.partition("|")
        r = p.add_run()
        r.text = "■  "
        style_run(r, size - 8, ACCENT)
        r = p.add_run()
        r.text = head
        style_run(r, size, WHITE, bool(rest))
        if rest:
            r = p.add_run()
            r.text = "  " + rest
            style_run(r, size, MUTED)
    return tb


def connect(slide, a, b, a_site=3, b_site=1, color=ACCENT, width=2):
    """Connection sites on rectangles: 0 top, 1 left, 2 bottom, 3 right."""
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
    c.begin_connect(a, a_site)
    c.end_connect(b, b_site)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd",
                          {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def link_run(paragraph, text, url, size=12, font=MONO):
    r = paragraph.add_run()
    r.text = text
    style_run(r, size, ACCENT, font=font)
    r.hyperlink.address = url
    return r


def set_cell_border(cell, color="2A2A38", width=12700):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
        ln = tcPr.makeelement(qn(tag), {"w": str(width)})
        fill = ln.makeelement(qn("a:solidFill"), {})
        fill.append(fill.makeelement(qn("a:srgbClr"), {"val": color}))
        ln.append(fill)
        tcPr.insert(0, ln)


def set_theme_link_colors(prs, hex_color="8B5CF6"):
    """Theme hlink/folHlink default to dark blue, unreadable on a dark background."""
    theme = prs.slide_master.part.part_related_by(RT.THEME)
    xml = theme.blob.decode("utf-8")
    for tag in ("hlink", "folHlink"):
        xml = re.sub(rf"<a:{tag}>.*?</a:{tag}>", f'<a:{tag}><a:srgbClr val="{hex_color}"/></a:{tag}>', xml, flags=re.S)
    theme._blob = xml.encode("utf-8")


# ================================================================ 1. Title
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = BG
rect(s, Inches(0.8), Inches(2.1), Inches(0.12), Inches(2.4), fill=ACCENT, line=None, shape=MSO_SHAPE.RECTANGLE)
t = textbox(s, Inches(1.15), Inches(1.95), Inches(11), Inches(1.2), "PersonaLab", 66, WHITE, True)
t.name = "Title"
textbox(s, Inches(1.15), Inches(3.15), Inches(11), Inches(0.8),
        "A Sokosumi Coworker that launches and grows an original AI influencer", 26, MUTED)
textbox(s, Inches(1.15), Inches(4.0), Inches(11), Inches(0.5),
        "Create  ·  Script  ·  Trends  ·  Analyze", 22, ACCENT, True)
textbox(s, Inches(1.15), Inches(5.6), Inches(11.5), Inches(0.9), [
    "Built with Masumi + Sokosumi on Cardano Preprod  ·  Vendor CardanoFish  ·  TOKEN2049 Origins Hackathon",
    f"Coworker {EVIDENCE['coworker_id']}",
], 14, MUTED)
s.notes_slide.notes_text_frame.text = (
    "[0:00-0:10] PersonaLab is a Sokosumi Coworker that launches and grows an original AI influencer. "
    "Four ready-to-run tasks: Create, Script, Trends and Analyze, each paid per task through Masumi on Cardano."
)

# ================================================================ 2. Problem
s = new_slide("AI-influencer tools are closed subscriptions", "Problem", notes="""
Creators and small brands pay a monthly subscription for a closed dashboard, even when they only need one job done.
Trend research is the slowest part and is usually guesswork without sources.
Nothing lets an agent marketplace sell these jobs one task at a time, with on-chain escrow and a verifiable result.
""")
bullets(s, Inches(0.85), Inches(1.75), Inches(11.5), Inches(4.8), [
    "Closed SaaS|$19-99 a month, credit-based, locked to one dashboard",
    "Four separate jobs|persona, scripts, trend research and analytics",
    "Trend research is guesswork|creators copy what they see, with no sources",
    "Made-up numbers|generic AI tools invent stats and follower counts",
    "No pay-per-task option|no marketplace sells these jobs per task with on-chain payment",
], 22, 16)

# ================================================================ 3. Solution
s = new_slide("One Coworker, four ready-to-run tasks", "Solution", notes="""
One Coworker with four task cards in Sokosumi. Each card's prompt carries a stage tag that the LangGraph supervisor routes on.
Create builds the persona, photos with an image model, a reel script and a 7-day plan.
Script writes trend-based Reels and TikTok scripts. Trends runs a swarm of parallel research agents and cites its evidence.
Analyze turns pasted metrics into KPIs, hook scores and next week's plan.
Schedule and Engage exist in code in dry-run mode but are not offered, because Sokosumi cannot link Instagram or TikTok accounts yet.
""")
cards = [
    ("Create", "AI persona card", "Persona photos via image model", "8-second reel script", "7-day post plan"),
    ("Script", "Trend-based scripts", "Reels and TikTok formats", "Hook, beats, shot list", "Caption + hashtags"),
    ("Trends", "Swarm of parallel research agents", "Hot short-video trends", "Cited evidence per trend", "Why it fits the persona"),
    ("Analyze", "Metrics in", "KPI summary (inputs shown)", "Hook scorecard", "Next week's plan"),
]
cw, ch, gap = Inches(2.85), Inches(3.5), Inches(0.2)
x0 = Inches(0.85)
for i, (name, *lines) in enumerate(cards):
    x = x0 + i * (cw + gap)
    c = rect(s, x, Inches(1.75), cw, ch, shape=MSO_SHAPE.RECTANGLE)
    rect(s, x, Inches(1.75), cw, Inches(0.09), fill=ACCENT, line=None, shape=MSO_SHAPE.RECTANGLE)
    textbox(s, x + Inches(0.2), Inches(2.0), cw - Inches(0.4), Inches(0.6), name, 26, WHITE, True)
    tb = textbox(s, x + Inches(0.2), Inches(2.75), cw - Inches(0.4), Inches(2.6), lines, 15, MUTED)
    for p in tb.text_frame.paragraphs:
        p.space_after = Pt(10)
textbox(s, Inches(0.85), Inches(5.55), Inches(11.6), Inches(0.8),
        "Roadmap, not offered: Schedule and Engage are in the code in dry-run mode. "
        "Sokosumi cannot link Instagram or TikTok accounts yet.", 14, MUTED)

# ================================================================ 4. Demo video
s = new_slide("Demo video", "Demo  ·  90 seconds", notes=f"""
HOW TO EMBED: select the placeholder box, delete it, then Insert > Video > This Device and place the recording in the same area.
The hackathon requires the video embedded in the file (no external links, no live demo).

90-SECOND TALK TRACK
[0:00-0:10] Title, then the Sokosumi Coworker page with the four cards.
  "PersonaLab is a Sokosumi Coworker that runs an AI influencer: Create, Script, Trends, Analyze."
[0:10-0:25] Open the Trends card and fill in the niche.
  "I pick Trends and give it my niche. The prompt starts with a stage tag."
[0:25-0:35] Worker log shows runtime start; the Task shows RUNNING.
  "Our Node worker picks the Task up and starts it with the Coworker's own runtime key."
[0:35-0:50] Graph log: the supervisor routes to the trend swarm; several researchers run in parallel.
  "A LangGraph supervisor reads the tag and fans out a swarm of research agents on DeepSeek V4.1 Flash."
[0:50-1:05] Task COMPLETED. Scroll the result: ranked trends, each with cited sources, then a Script run that uses them.
  "Every trend comes with its evidence. No invented numbers."
[1:05-1:20] Payment: MPS terms, escrow FundsLocked on preprod.cardanoscan.io, result hash, collection (when done).
  "Paid Tasks settle through Masumi: 1 test USDM is locked in escrow, the result hash is submitted, and the seller collects on Cardano Preprod."
[1:20-1:30] Evidence slide.
  "Every step is journaled, so nothing runs or pays twice. PersonaLab: your AI influencer team, hired by the task."

Evidence Tasks: {EVIDENCE['task_personal']} (personal), {EVIDENCE['task_event']} (event workspace), paid {EVIDENCE['task_paid']}.
Recording: 1080p; hide .env.local, terminal history, seed phrases, wallet output, email addresses and local paths.
""")
ph = rect(s, Inches(0.85), Inches(1.6), Inches(7.6), Inches(4.3), fill=ACCENT_SOFT, line=ACCENT, radius=0.04)
ph.line.dash_style = MSO_LINE_DASH_STYLE.DASH
ph.name = "Demo video placeholder"
box_text(ph, ["Demo video goes here",
              "Insert > Video > This Device",
              "Delete this box and place the 90-second recording here (16:9).",
              "Must be embedded in the file, not linked."], 22, first_color=WHITE)
steps = [
    "1. Pick a card in Sokosumi, fill in the prompt",
    "2. Task READY, worker runs runtime start",
    "3. Supervisor routes on the [stage:*] tag",
    "4. Agents run; result saved to the journal",
    "5. runtime complete, Task COMPLETED",
    "6. Paid path: escrow, result hash, collection",
]
tb = textbox(s, Inches(8.8), Inches(1.6), Inches(3.9), Inches(4.3), ["Flow"] + steps, 15, MUTED)
style_run(tb.text_frame.paragraphs[0].runs[0], 18, ACCENT, True)
for p in tb.text_frame.paragraphs:
    p.space_after = Pt(12)

# ================================================================ 5. Architecture
s = new_slide("Architecture", "How it runs", notes="""
A Sokosumi Task is polled by our Node worker (one worker lock, a journal per Task). The worker calls runtime start and runtime complete
with the Coworker key from the OS vault, and handles Masumi payment through the Masumi Payment Service on Cardano Preprod.
The worker calls the Python LangGraph service over loopback HTTP with a bearer token. The supervisor routes to the stage's agents.
Text work runs on DeepSeek V4.1 Flash, with OpenRouter as fallback. Persona photos use an image model.
""")
Y = Inches(1.9)
H = Inches(1.15)
sok = rect(s, Inches(0.7), Y, Inches(2.5), H); box_text(sok, ["Sokosumi Task", "card prompt [stage:*]"], 16)
wrk = rect(s, Inches(3.9), Y, Inches(2.8), H); box_text(wrk, ["Node worker", "Sokosumi runtime + Masumi payment"], 16)
sup = rect(s, Inches(7.4), Y, Inches(2.6), H, fill=ACCENT_SOFT, line=ACCENT)
box_text(sup, ["LangGraph supervisor", "Python, FastAPI, loopback"], 16)
connect(s, sok, wrk)
connect(s, wrk, sup)

ag_x = Inches(10.65)
ag_w, ag_h = Inches(2.1), Inches(0.72)
agent_boxes = []
for i, (n, sub) in enumerate([("Script", "script agent"), ("Trends", "research swarm"),
                              ("Analyze", "analyst agent"), ("Create", "persona, content, media")]):
    b = rect(s, ag_x, Inches(1.35) + i * Inches(0.85), ag_w, ag_h)
    box_text(b, [n, sub], 14)
    agent_boxes.append(b)
    connect(s, sup, b, 3, 1, width=1.25)

mps = rect(s, Inches(3.9), Inches(4.4), Inches(2.8), Inches(1.0))
box_text(mps, ["Masumi Payment Service", "terms, escrow, result hash"], 14)
ada = rect(s, Inches(0.7), Inches(4.4), Inches(2.5), Inches(1.0))
box_text(ada, ["Cardano Preprod", "Web3CardanoV2 escrow"], 14)
connect(s, wrk, mps, 2, 0)
connect(s, mps, ada, 1, 3)

llm = rect(s, Inches(7.4), Inches(4.4), Inches(2.6), Inches(1.0))
box_text(llm, ["DeepSeek V4.1 Flash", "OpenRouter fallback"], 14)
img = rect(s, Inches(10.65), Inches(4.85), Inches(2.1), Inches(0.9))
box_text(img, ["Image model", "Create: persona photos"], 14)
connect(s, sup, llm, 2, 0)
connect(s, agent_boxes[-1], img, 2, 0, width=1.25)

textbox(s, Inches(0.7), Inches(5.95), Inches(12), Inches(0.6),
        "Masumi agent registered: RegistrationConfirmed  ·  Worker polls every 5 s, one lock, journal per Task  ·  "
        "Secrets in .env.local and the OS vault only", 13, MUTED)

# ================================================================ 6. LangGraph agents + swarm
s = new_slide("LangGraph agents and the trend swarm", "Inside the graph", notes="""
The supervisor reads the stage tag from the card prompt; untagged input is classified by the model.
Create chains persona, content and media agents. Script and Analyze are single focused agents.
Trends is a swarm: the work is split into research angles, parallel research agents run at the same time,
and a merge step dedupes and ranks the trends. A trend without a cited source is dropped.
""")
# left: routing list
tb = textbox(s, Inches(0.85), Inches(1.7), Inches(5.6), Inches(4.6), ["Supervisor routing"], 18, ACCENT, True)
bullets(s, Inches(0.85), Inches(2.3), Inches(5.6), Inches(4.0), [
    "[stage:create]|persona → content → media (image model)",
    "[stage:script]|trend-based Reels and TikTok scripts",
    "[stage:trends]|fan-out to the research swarm",
    "[stage:analyze]|KPIs, hook scores, next week",
    "No tag|the model classifies the request",
], 16, 10)
# right: swarm diagram
plan = rect(s, Inches(7.0), Inches(1.75), Inches(5.6), Inches(0.75), fill=ACCENT_SOFT, line=ACCENT)
box_text(plan, ["Split niche into research angles"], 15, bold_first=True)
rs = []
for i in range(4):
    r = rect(s, Inches(7.0) + i * Inches(1.45), Inches(3.05), Inches(1.25), Inches(0.95))
    box_text(r, [f"Agent {i + 1}", "researcher"], 13)
    rs.append(r)
    connect(s, plan, r, 2, 0, width=1.25)
merge = rect(s, Inches(7.0), Inches(4.55), Inches(5.6), Inches(0.75), fill=ACCENT_SOFT, line=ACCENT)
box_text(merge, ["Merge: dedupe, rank, keep only cited trends"], 15)
for r in rs:
    connect(s, r, merge, 2, 0, width=1.25)
textbox(s, Inches(7.0), Inches(5.5), Inches(5.6), Inches(0.8),
        "Output: hot short-video trends, each with source links and why it fits the persona", 14, MUTED)

# ================================================================ 7. Masumi payment flow
s = new_slide("Masumi payment flow", "Paid per task on Cardano Preprod", notes=f"""
1. Terms: the buyer starts a paid Task; MPS quotes 1 tUSDM with signed deadlines (payByTime, submitResultTime, unlockTime).
2. Escrow: the buyer's masumiPayment locks funds in the Web3CardanoV2 contract; the worker waits for on-chain confirmation.
   Paid Task {EVIDENCE['task_paid']} reached FundsLocked.
3. Result hash: the worker runs the graph, submits the SHA-256 result hash to MPS and completes the Task.
4. Collection: after unlockTime the seller collects. settlement.mjs checks Core's receipt against the MPS withdrawal tx and
   Blockfrost UTxOs and reports the net amount received. Collection is still pending; do not claim it until the tx is confirmed.
Every step is saved before the next outside write, so a crash never pays or collects twice.
""")
steps = [
    ("1  Terms", "1 tUSDM quote", "signed deadlines", "Done"),
    ("2  Escrow", "funds locked in", "Web3CardanoV2", "FundsLocked"),
    ("3  Result hash", "SHA-256 of result", "submitted to MPS", "TODO: verify"),
    ("4  Collection", "seller withdraws", "after unlockTime", "Pending"),
]
sw_, gap = Inches(2.75), Inches(0.37)
prev = None
for i, (head, a, b, status) in enumerate(steps):
    x = Inches(0.85) + i * (sw_ + gap)
    done = status in ("Done", "FundsLocked")
    c = rect(s, x, Inches(1.9), sw_, Inches(2.3), fill=ACCENT_SOFT if done else CARD, line=ACCENT if done else CARD_LINE)
    box_text(c, [head, a, b], 20)
    chip = rect(s, x + Inches(0.45), Inches(4.4), sw_ - Inches(0.9), Inches(0.5),
                fill=ACCENT if done else CARD, line=None if done else MUTED, radius=0.5)
    box_text(chip, [status], 13, bold_first=True)
    if prev is not None:
        connect(s, prev, c)
    prev = c
bullets(s, Inches(0.85), Inches(5.2), Inches(11.8), Inches(1.4), [
    "Crash-safe|each step is journaled before the next outside write: never pays or collects twice",
    "Verified receipt|settlement.mjs cross-checks Core, the MPS withdrawal tx and Blockfrost UTxOs",
], 15, 6)

# ================================================================ 8. Evidence
s = new_slide("Evidence", "Verified on Sokosumi and Cardano Preprod", notes=f"""
All IDs here were read back from Sokosumi, MPS or the chain. TODO cells are not done yet and must not be filled with guesses.
Full Masumi agentIdentifier: {EVIDENCE['agent_identifier']}
Registration tx: {EXPLORER}{EVIDENCE['registration_tx']}
Update the EVIDENCE dict in docs/deck/build_deck.py and re-run it when the collection tx is confirmed.
""")
ai = EVIDENCE["agent_identifier"]
rows = [
    ("Coworker", EVIDENCE["coworker_id"], None),
    ("Masumi agent", f"RegistrationConfirmed  ·  {ai[:16]}…{ai[-12:]}", None),
    ("Registration tx", EVIDENCE["registration_tx"], EVIDENCE["registration_tx"]),
    ("Task, personal (Engage)", f"{EVIDENCE['task_personal']}  ·  COMPLETED", None),
    ("Task, event workspace", f"{EVIDENCE['task_event']}  ·  COMPLETED", None),
    ("Paid Task", f"{EVIDENCE['task_paid']}  ·  {EVIDENCE['paid_status']}", None),
    ("Escrow tx", EVIDENCE["escrow_tx"], None if is_todo(EVIDENCE["escrow_tx"]) else EVIDENCE["escrow_tx"]),
    ("Collection tx", EVIDENCE["collection_tx"],
     None if is_todo(EVIDENCE["collection_tx"]) else EVIDENCE["collection_tx"]),
    ("Net seller receipt", EVIDENCE["net_receipt"], None),
]
tbl_shape = s.shapes.add_table(len(rows) + 1, 2, Inches(0.85), Inches(1.6), Inches(11.6), Inches(4.9))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(3.0)
tbl.columns[1].width = Inches(8.6)
tbl.first_row = True
for r_i, (k, v, tx) in enumerate([("Item", "Value", None)] + rows):
    for c_i, txt in enumerate((k, v)):
        cell = tbl.cell(r_i, c_i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT_SOFT if r_i == 0 else (CARD if r_i % 2 else BG)
        cell.margin_left = Inches(0.12)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        set_cell_border(cell)
        tf = cell.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        if c_i == 1 and tx:
            link_run(p, txt, EXPLORER + tx, 12)
        else:
            r = p.add_run()
            r.text = txt
            todo = is_todo(txt)
            style_run(r, 13 if c_i == 0 else 12, ACCENT if todo else (WHITE if c_i == 0 or r_i == 0 else MUTED),
                      r_i == 0 or c_i == 0, FONT if (c_i == 0 or r_i == 0) else MONO)
    tbl.rows[r_i].height = Inches(0.48)

# ================================================================ 9. Quality checks
s = new_slide("Quality and safety checks", "Trust", notes="""
Every card has fixed section headings, and the smoke test counts them.
The prompts forbid invented statistics, follower counts and testimonials; unsure figures are marked [verify].
Trends keeps only trends with a cited source. Personas are always original, with AI disclosure recommended.
Pasted comments, metrics and web pages are treated as untrusted data. Analyze returns a measurement plan if no metrics are given.
""")
bullets(s, Inches(0.85), Inches(1.75), Inches(11.6), Inches(4.8), [
    "Fixed output sections|the smoke test counts every ## heading",
    "No invented numbers|unsure figures are marked [verify]",
    "Cited trends only|a trend without a source link is dropped",
    "Original personas|never a real person's likeness; AI disclosure recommended",
    "Untrusted input|pasted text and web pages are data, not instructions",
    "Ops|bearer-token loopback graph, secrets only in .env.local and the OS vault, errors logged without prompts or keys",
], 19, 14)

# ================================================================ 10. Human help
s = new_slide("What needed human help", "Honest log", notes="""
These steps could not be done by the agent alone, and are recorded in docs/setup-state.md with the error history.
The biggest scope change: Sokosumi cannot link Instagram or TikTok accounts, so Schedule and Engage were taken off the card list.
""")
bullets(s, Inches(0.85), Inches(1.75), Inches(11.6), Inches(4.8), [
    "Login|OAuth in the owner's browser, plus a fix for Windows cutting the URL at &",
    "Access|reusing the Vendor and Coworker; organizer approval for the event Workspace",
    "Keys and funds|model API key, Blockfrost Preprod key, test ADA and USDM from the faucet",
    "Local infra|starting Docker Desktop for the MPS Postgres database",
    "Payment|approving the paid Task; seller collection still pending",
    "Scope|no Instagram/TikTok linking in Sokosumi, so Schedule and Engage moved to the roadmap",
], 19, 14)

# ================================================================ 11. Roadmap
s = new_slide("Roadmap", "Next", notes="""
Schedule and Engage are already written and tested in dry-run mode; they go live when Sokosumi supports linking social accounts.
After that: video reels and voice, human approval through LangGraph interrupts, each stage as its own hireable Masumi agent, and a hosted deployment.
""")
bullets(s, Inches(0.85), Inches(1.75), Inches(11.6), Inches(4.8), [
    "Schedule + Engage|in code, dry-run today; go live once Sokosumi can link Instagram/TikTok",
    "Video|image-to-video reels and voiceover from the persona photos",
    "Human approval|LangGraph interrupt() with a Postgres checkpointer, approved via Task comments",
    "Agent-to-agent|each stage as its own Masumi agent, hireable through /v1/tasks/{id}/jobs",
    "Hosted|graph, worker and MPS deployed; pass a laptop-offline test",
], 19, 14)

# ================================================================ 12. Team / links
s = new_slide("Team and links", "Contact", notes="""
The repository is private. Judges need to be given access before review: share the GitHub handles with the team.
""")
lines = [
    ("Repository", EVIDENCE["repo"], EVIDENCE["repo"]),
    ("Repo access", "Private: judges need access (send GitHub handles)", None),
    ("Coworker", EVIDENCE["coworker_id"], None),
    ("Vendor", EVIDENCE["vendor"], None),
    ("Masumi registration", EVIDENCE["registration_tx"][:24] + "…", EXPLORER + EVIDENCE["registration_tx"]),
    ("Team", EVIDENCE["team"], None),
]
tb = s.shapes.add_textbox(Inches(0.85), Inches(1.75), Inches(11.6), Inches(4.6))
tf = tb.text_frame
tf.word_wrap = True
for i, (k, v, url) in enumerate(lines):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(16)
    r = p.add_run()
    r.text = f"{k}:  "
    style_run(r, 20, WHITE, True)
    if url:
        link_run(p, v, url, 18, FONT)
    else:
        r = p.add_run()
        r.text = v
        style_run(r, 18, ACCENT if is_todo(v) or k == "Repo access" else MUTED)
textbox(s, Inches(0.85), Inches(6.2), Inches(11.6), Inches(0.5),
        "PersonaLab: your AI influencer team, hired by the task.", 20, ACCENT, True)

set_theme_link_colors(prs)
prs.save(OUT)
print(f"saved {OUT}")
